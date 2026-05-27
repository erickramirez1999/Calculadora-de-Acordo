"""
Importador de XLSX de títulos (briefing Seção 17).

Colunas obrigatórias (cabeçalho na linha 1):
  - Código do Parceiro
  - Razão Social do Parceiro
  - Empresa (1 ou 2)
  - Código do Vendedor
  - Nome do Vendedor
  - Data de Vencimento (dd/mm/aaaa)
  - Nº da Nota
  - Nº Único
  - Valor Principal (numérico R$)

O atraso em dias NÃO vem no arquivo — é calculado pelo sistema na hora.

Validações:
  - data válida
  - principal > 0
  - Nº Único não duplicado dentro do arquivo
  - Código do Parceiro preenchido
  - Empresa ∈ {1, 2}
  - Código do Vendedor preenchido

Aceita .xls (engine xlrd) e .xlsx (engine openpyxl).

NOTA SOBRE COMPATIBILIDADE: o leitor antigo (BASE_CRUA do Sankhya) tinha um formato
diferente. Mantemos um detector flexível: se as colunas do briefing forem encontradas,
usa esse formato; senão, tenta o formato BASE_CRUA antigo. Permite o usuário usar
qualquer um dos dois.
"""
from __future__ import annotations

import io
from datetime import date, datetime
from typing import List, Optional, Tuple

import pandas as pd

from src.modelos.tipos import Boleto, Empresa


# ============================================================
# COLUNAS ESPERADAS - FORMATO BRIEFING (Seção 17)
# ============================================================

COLUNAS_BRIEFING = {
    "Código do Parceiro": "codigo_parceiro",
    "Razão Social do Parceiro": "razao_social_parceiro",
    "Empresa": "empresa",
    "Código do Vendedor": "codigo_vendedor",
    "Nome do Vendedor": "nome_vendedor",
    "Data de Vencimento": "vencimento",
    "Nº da Nota": "numero_nota",
    "Nº Único": "numero_unico",
    "Valor Principal": "principal",
}

# Aliases pra lidar com variações: formato do briefing E formato BASE_CRUA do Sankhya
ALIASES = {
    # === Código do Parceiro ===
    "codigo do parceiro": "codigo_parceiro",
    "cod parceiro": "codigo_parceiro",
    "cód parceiro": "codigo_parceiro",
    "cod. parceiro": "codigo_parceiro",
    "parceiro": "codigo_parceiro",                      # BASE_CRUA
    # === Razão Social ===
    "razao social do parceiro": "razao_social_parceiro",
    "razão social do parceiro": "razao_social_parceiro",
    "razao social": "razao_social_parceiro",
    "razão social": "razao_social_parceiro",
    "nome parceiro (parceiro)": "razao_social_parceiro",  # BASE_CRUA
    "nome do parceiro": "razao_social_parceiro",
    "nome parceiro": "razao_social_parceiro",
    # === Empresa ===
    "empresa": "empresa",
    # === Código do Vendedor ===
    "codigo do vendedor": "codigo_vendedor",
    "código do vendedor": "codigo_vendedor",
    "cod vendedor": "codigo_vendedor",
    "cód vendedor": "codigo_vendedor",
    "vendedor": "codigo_vendedor",                      # BASE_CRUA
    # === Nome do Vendedor ===
    "nome do vendedor": "nome_vendedor",
    "nome vendedor": "nome_vendedor",
    "apelido": "nome_vendedor",                         # BASE_CRUA
    "apelido (vendedor)": "nome_vendedor",
    # === Vencimento ===
    "data de vencimento": "vencimento",
    "vencimento": "vencimento",
    "dt. vencimento": "vencimento",                     # BASE_CRUA
    "dt vencimento": "vencimento",
    "data vencimento": "vencimento",
    # === Nº da Nota ===
    "n da nota": "numero_nota",
    "nº da nota": "numero_nota",
    "no da nota": "numero_nota",
    "nro nota": "numero_nota",                          # BASE_CRUA
    "numero nota": "numero_nota",
    "número nota": "numero_nota",
    "nf": "numero_nota",
    # === Nº Único ===
    "n único": "numero_unico",
    "nº único": "numero_unico",
    "no unico": "numero_unico",
    "n unico": "numero_unico",
    "nro único": "numero_unico",                        # BASE_CRUA
    "nro unico": "numero_unico",
    "numero unico": "numero_unico",
    "número único": "numero_unico",
    # === Valor Principal ===
    "valor principal": "principal",
    "vlr do desdobramento": "principal",                # BASE_CRUA
    "vlr desdobramento": "principal",
    "valor": "principal",
    "principal": "principal",
    "valor liquido": "principal",
    "valor líquido": "principal",
    # === Tipo de Título (filtro de boletos) ===
    # No Sankhya, a coluna "Tipo de Título" tem código numérico.
    # Só importamos os tipos que são BOLETO/COBRANÇA. Os outros
    # (crédito automático, depósito bancário, etc) são ignorados.
    "tipo de titulo": "tipo_titulo",
    "tipo de título": "tipo_titulo",
    "tipo titulo": "tipo_titulo",
    "tipo título": "tipo_titulo",
}


# Tipos de título permitidos (decisão Erick - 13/05/2026):
# Esses são os códigos do Sankhya que representam BOLETOS / formas de cobrança
# tradicionais que o sistema gerencia. Outros tipos (crédito automático,
# depósito bancário, transferência etc) são ignorados na importação.
TIPOS_TITULO_PERMITIDOS = {4, 28, 29, 39, 40, 41, 47, 48, 64, 70}


def _normalizar_nome_coluna(nome: str) -> str:
    """Normaliza nome de coluna pra match com aliases (lowercase + sem acentos)."""
    import unicodedata
    s = unicodedata.normalize("NFKD", str(nome))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return s.lower().strip()


def _detectar_engine(nome_arquivo: str) -> str:
    return "xlrd" if nome_arquivo.lower().endswith(".xls") else "openpyxl"


def _achar_linha_cabecalho(arquivo_bytes: bytes, engine: str) -> int:
    """
    Procura nas primeiras 15 linhas qual contém termos típicos de cabeçalho.
    Cobre tanto o formato do briefing quanto o BASE_CRUA do Sankhya.
    """
    bio = io.BytesIO(arquivo_bytes)
    preview = pd.read_excel(bio, header=None, nrows=15, engine=engine)
    palavras_chave_header = {
        "data de vencimento", "dt. vencimento", "vencimento", "dt vencimento",
        "nro único", "nro unico", "nº único", "numero unico",
        "nro nota", "nº da nota", "numero nota",
    }
    for i, row in preview.iterrows():
        normalizado = {_normalizar_nome_coluna(v) for v in row.values if pd.notna(v)}
        if normalizado & palavras_chave_header:
            return int(i)
    raise ValueError(
        "Não encontrei a linha do cabeçalho. "
        "Esperava encontrar 'Data de Vencimento' (ou 'Dt. Vencimento') nas 15 primeiras linhas."
    )


def _mapear_colunas(df: pd.DataFrame) -> dict:
    """Cria um dict mapeando nome interno → nome real no df."""
    mapa = {}
    for col_original in df.columns:
        normalizado = _normalizar_nome_coluna(col_original)
        if normalizado in ALIASES:
            mapa[ALIASES[normalizado]] = col_original
    return mapa


def _parsear_data(v) -> Optional[date]:
    """Tenta converter vários formatos pra date."""
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        s = v.strip()
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y", "%Y-%m-%d %H:%M:%S"):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                continue
    try:
        return pd.to_datetime(v, dayfirst=True).date()
    except Exception:
        return None


def _parsear_valor(v) -> Optional[float]:
    """Aceita 1234.56, 1234,56, '1.234,56' etc."""
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("R$", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _parsear_empresa(v) -> Optional[Empresa]:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    try:
        i = int(float(str(v).strip()))
        return Empresa(i)
    except (ValueError, KeyError):
        return None


def _limpar_codigo(v) -> str:
    """
    Limpa códigos numéricos/texto que vêm das planilhas.
    Remove .0 do final quando o valor é float-inteiro (37299.0 -> 37299).
    Remove espaços e trata None/NaN.
    """
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    s = str(v).strip()
    if not s or s.lower() == "nan":
        return ""
    # Se é um número float inteiro tipo "37299.0", remove o ".0"
    try:
        f = float(s)
        if f.is_integer():
            return str(int(f))
    except ValueError:
        pass
    return s


# ============================================================
# RESULTADO DA IMPORTAÇÃO
# ============================================================

class ResultadoImportacao:
    def __init__(self):
        self.boletos: List[Boleto] = []
        self.erros: List[str] = []          # erros bloqueantes (arquivo todo inválido)
        self.avisos_linhas: List[Tuple[int, str]] = []   # (linha, msg) — só essas linhas
        self.avisos_gerais: List[str] = []  # avisos do arquivo todo (ex: tipos filtrados)
        self.total_linhas: int = 0

    @property
    def tem_erros_bloqueantes(self) -> bool:
        return len(self.erros) > 0

    @property
    def qtd_validas(self) -> int:
        return len(self.boletos)

    @property
    def qtd_invalidas(self) -> int:
        return len(self.avisos_linhas)


# ============================================================
# FUNÇÃO PRINCIPAL
# ============================================================

def importar_xlsx_titulos(arquivo_bytes: bytes, nome_arquivo: str) -> ResultadoImportacao:
    """
    Lê um arquivo XLSX/XLS com títulos no formato do briefing (Seção 17).
    Retorna ResultadoImportacao com boletos válidos + lista de avisos.
    """
    res = ResultadoImportacao()

    try:
        engine = _detectar_engine(nome_arquivo)
        linha_cab = _achar_linha_cabecalho(arquivo_bytes, engine)
        bio = io.BytesIO(arquivo_bytes)
        df = pd.read_excel(bio, header=linha_cab, engine=engine)
    except Exception as e:
        res.erros.append(f"Falha ao abrir o arquivo: {e}")
        return res

    df.columns = [str(c).strip() for c in df.columns]
    mapa = _mapear_colunas(df)

    obrigatorias = [
        "codigo_parceiro", "razao_social_parceiro", "empresa",
        "codigo_vendedor", "nome_vendedor",
        "vencimento", "numero_nota", "numero_unico", "principal",
    ]
    faltantes = [c for c in obrigatorias if c not in mapa]
    if faltantes:
        nomes_amigaveis = {
            "codigo_parceiro": "Código do Parceiro",
            "razao_social_parceiro": "Razão Social do Parceiro",
            "empresa": "Empresa",
            "codigo_vendedor": "Código do Vendedor",
            "nome_vendedor": "Nome do Vendedor",
            "vencimento": "Data de Vencimento",
            "numero_nota": "Nº da Nota",
            "numero_unico": "Nº Único",
            "principal": "Valor Principal",
        }
        res.erros.append(
            "Colunas obrigatórias faltando no arquivo: "
            + ", ".join(nomes_amigaveis[c] for c in faltantes)
        )
        return res

    numeros_unicos_vistos = set()
    res.total_linhas = len(df)

    # Contadores extras pro filtro de tipo
    qtd_filtrados_tipo = 0
    tipos_ignorados = set()

    for idx, row in df.iterrows():
        linha_real = idx + linha_cab + 2  # +2 porque header é 0-indexed e linha 1 do Excel é título

        # Pula linhas totalmente em branco
        if all(pd.isna(row[mapa[c]]) for c in obrigatorias):
            continue

        # FILTRO DE TIPO DE TÍTULO (Erick - 13/05/2026):
        # Se a planilha tem a coluna "Tipo de Título" e o tipo NÃO está
        # na lista de tipos permitidos (boletos), o registro é silenciosamente
        # ignorado. Outros tipos (crédito automático, depósito, etc) não devem
        # virar acordo.
        if "tipo_titulo" in mapa:
            tipo_raw = row[mapa["tipo_titulo"]]
            if not pd.isna(tipo_raw):
                try:
                    tipo_int = int(float(str(tipo_raw)))
                    if tipo_int not in TIPOS_TITULO_PERMITIDOS:
                        qtd_filtrados_tipo += 1
                        tipos_ignorados.add(tipo_int)
                        continue
                except (ValueError, TypeError):
                    # Se não conseguiu converter, deixa passar pra
                    # demais validações tratarem
                    pass

        # Validações campo a campo
        venc = _parsear_data(row[mapa["vencimento"]])
        if venc is None:
            res.avisos_linhas.append((linha_real, "Data de Vencimento inválida ou vazia"))
            continue

        principal = _parsear_valor(row[mapa["principal"]])
        if principal is None or principal <= 0:
            res.avisos_linhas.append((linha_real, "Valor Principal precisa ser > 0"))
            continue

        empresa = _parsear_empresa(row[mapa["empresa"]])
        if empresa is None:
            res.avisos_linhas.append((linha_real, "Empresa precisa ser 1 ou 2"))
            continue

        codigo_parceiro = _limpar_codigo(row[mapa["codigo_parceiro"]])
        if not codigo_parceiro:
            res.avisos_linhas.append((linha_real, "Código do Parceiro vazio"))
            continue

        codigo_vendedor = _limpar_codigo(row[mapa["codigo_vendedor"]])
        if not codigo_vendedor:
            res.avisos_linhas.append((linha_real, "Código do Vendedor vazio"))
            continue

        nro_unico = _limpar_codigo(row[mapa["numero_unico"]])
        if not nro_unico:
            res.avisos_linhas.append((linha_real, "Nº Único vazio"))
            continue
        if nro_unico in numeros_unicos_vistos:
            res.avisos_linhas.append(
                (linha_real, f"Nº Único duplicado dentro do arquivo: {nro_unico}")
            )
            continue
        numeros_unicos_vistos.add(nro_unico)

        nro_nota = _limpar_codigo(row[mapa["numero_nota"]])
        razao = str(row[mapa["razao_social_parceiro"]] or "").strip()
        nome_vendedor = str(row[mapa["nome_vendedor"]] or "").strip()

        boleto = Boleto(
            codigo_parceiro=codigo_parceiro,
            razao_social_parceiro=razao,
            empresa=empresa,
            codigo_vendedor=codigo_vendedor,
            nome_vendedor=nome_vendedor,
            vencimento=venc,
            numero_nota=nro_nota,
            numero_unico=nro_unico,
            principal=round(principal, 2),
        )
        res.boletos.append(boleto)

    # Se filtrou títulos por tipo, adiciona um aviso amigável
    if qtd_filtrados_tipo > 0:
        tipos_str = ", ".join(str(t) for t in sorted(tipos_ignorados))
        res.avisos_gerais.append(
            f"ℹ️ {qtd_filtrados_tipo} título(s) ignorado(s) por não serem boleto/cobrança "
            f"(tipos {tipos_str}). Só importamos: 4, 28, 29, 39, 40, 41, 47, 48, 64, 70."
        )

    return res


def montar_relatorio_importacao(res: ResultadoImportacao) -> str:
    """Texto amigável pra mostrar no Streamlit após import."""
    linhas = []
    if res.tem_erros_bloqueantes:
        linhas.append("⚠️ Não foi possível importar o arquivo:")
        for e in res.erros:
            linhas.append(f"   • {e}")
        return "\n".join(linhas)

    linhas.append(f"📥 Arquivo lido: {res.total_linhas} linhas no total.")
    linhas.append(f"✅ Importados com sucesso: {res.qtd_validas} títulos.")
    if res.qtd_invalidas:
        linhas.append(f"⚠️ Ignorados por erro: {res.qtd_invalidas} linhas.")
        for ln, msg in res.avisos_linhas[:10]:
            linhas.append(f"   - Linha {ln}: {msg}")
        if len(res.avisos_linhas) > 10:
            linhas.append(f"   ... e mais {len(res.avisos_linhas) - 10} avisos.")
    return "\n".join(linhas)


# ============================================================
# TEMPLATE XLSX vazio para o usuário baixar
# ============================================================

def gerar_template_xlsx() -> bytes:
    """Gera um XLSX modelo com cabeçalhos e 1 linha de exemplo, pra usuário baixar."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Títulos"

    headers = list(COLUNAS_BRIEFING.keys())
    fonte_bold = Font(bold=True, color="FFFFFF", name="Montserrat")
    fill = PatternFill("solid", start_color="041747")
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    borda = Border(*([Side(style="thin", color="D9D9D9")] * 4))

    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = fonte_bold
        c.fill = fill
        c.alignment = align
        c.border = borda
        ws.column_dimensions[get_column_letter(i)].width = 22

    # Linha exemplo (genérica - sem citar clientes reais)
    ws.cell(row=2, column=1, value="37299")
    ws.cell(row=2, column=2, value="EMPRESA EXEMPLO LTDA")
    ws.cell(row=2, column=3, value=1)
    ws.cell(row=2, column=4, value="2235")
    ws.cell(row=2, column=5, value="NOME DO VENDEDOR")
    ws.cell(row=2, column=6, value=datetime(2026, 2, 19))
    ws.cell(row=2, column=6).number_format = "dd/mm/yyyy"
    ws.cell(row=2, column=7, value="2291660")
    ws.cell(row=2, column=8, value="16475907")
    ws.cell(row=2, column=9, value=1778.65)
    ws.cell(row=2, column=9).number_format = "#,##0.00"

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
